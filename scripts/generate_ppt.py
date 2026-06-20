#!/usr/bin/env python3
"""Generate defense PPT for word-memory project."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Colors
GREEN_DARK = RGBColor(0x1B, 0x43, 0x32)
GREEN_MID = RGBColor(0x2D, 0x6A, 0x4F)
GREEN_LIGHT = RGBColor(0x52, 0xB7, 0x88)
GREEN_PALE = RGBColor(0xD8, 0xF3, 0xDC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WARM_WHITE = RGBColor(0xFA, 0xF8, 0xF3)
TEXT_DARK = RGBColor(0x2C, 0x3E, 0x2D)
TEXT_MUTED = RGBColor(0x6B, 0x7F, 0x6E)


def set_slide_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_top_bar(slide, prs):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN_MID
    bar.line.fill.background()


def add_title(slide, text, top=Inches(0.55), size=32, color=GREEN_DARK):
    box = slide.shapes.add_textbox(Inches(0.7), top, Inches(11.9), Inches(0.9))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = color
    return box


def add_body(slide, lines, top=Inches(1.6), width=11.5, size=20, bullet=True):
    box = slide.shapes.add_textbox(Inches(0.7), top, Inches(width), Inches(5.5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " if bullet else "") + line
        p.font.size = Pt(size)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(10)
        p.level = 0
    return box


def add_footer(slide, text="词忆 · 单词记忆网站 | 答辩演示"):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(7.0), Inches(11), Inches(0.4))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED


def add_card(slide, left, top, width, height, title, items):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = GREEN_PALE
    shape.line.width = Pt(1.5)

    tbox = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), height - Inches(0.3))
    tf = tbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GREEN_MID
    for item in items:
        bp = tf.add_paragraph()
        bp.text = "• " + item
        bp.font.size = Pt(14)
        bp.font.color.rgb = TEXT_DARK
        bp.space_before = Pt(4)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Slide 1 - Cover
    s1 = prs.slides.add_slide(blank)
    set_slide_bg(s1, WARM_WHITE)
    add_top_bar(s1, prs)

    logo = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.9), Inches(1.5), Inches(1.5), Inches(1.5))
    logo.fill.solid()
    logo.fill.fore_color.rgb = GREEN_MID
    logo.line.fill.background()
    lt = logo.text_frame
    lt.paragraphs[0].text = "词"
    lt.paragraphs[0].font.size = Pt(48)
    lt.paragraphs[0].font.bold = True
    lt.paragraphs[0].font.color.rgb = WHITE
    lt.paragraphs[0].alignment = PP_ALIGN.CENTER
    lt.vertical_anchor = MSO_ANCHOR.MIDDLE

    title = s1.shapes.add_textbox(Inches(1), Inches(3.3), Inches(11.3), Inches(1))
    tp = title.text_frame.paragraphs[0]
    tp.text = "词忆 · 单词记忆网站"
    tp.font.size = Pt(44)
    tp.font.bold = True
    tp.font.color.rgb = GREEN_DARK
    tp.alignment = PP_ALIGN.CENTER

    sub = s1.shapes.add_textbox(Inches(1), Inches(4.3), Inches(11.3), Inches(0.8))
    sp = sub.text_frame.paragraphs[0]
    sp.text = "Spring Boot + 艾宾浩斯间隔复习 | 项目答辩"
    sp.font.size = Pt(22)
    sp.font.color.rgb = TEXT_MUTED
    sp.alignment = PP_ALIGN.CENTER

    info = s1.shapes.add_textbox(Inches(1), Inches(5.5), Inches(11.3), Inches(1))
    ip = info.text_frame.paragraphs[0]
    ip.text = "GitHub: github.com/Nancy985/word-memory\n演示账号: demo / password123"
    ip.font.size = Pt(16)
    ip.font.color.rgb = GREEN_MID
    ip.alignment = PP_ALIGN.CENTER

    # Slide 2 - Background
    s2 = prs.slides.add_slide(blank)
    set_slide_bg(s2, WARM_WHITE)
    add_top_bar(s2, prs)
    add_title(s2, "01  项目背景与目标")
    add_body(s2, [
        "背景：传统背单词缺乏科学复习计划，容易遗忘",
        "目标：开发一款 Web 单词记忆系统，帮助用户高效背单词",
        "用户：需要备考、日常积累词汇的学习者",
        "周期：2025年6月12日 — 6月21日",
        "交付：前后端互通、GitHub 开源、Apifox 接口文档",
    ])
    add_footer(s2)

    # Slide 3 - Features
    s3 = prs.slides.add_slide(blank)
    set_slide_bg(s3, WARM_WHITE)
    add_top_bar(s3, prs)
    add_title(s3, "02  功能模块")
    add_card(s3, Inches(0.6), Inches(1.5), Inches(3.8), Inches(2.2), "用户认证", ["注册 / 登录", "Token 鉴权", "BCrypt 加密"])
    add_card(s3, Inches(4.7), Inches(1.5), Inches(3.8), Inches(2.2), "单词书管理", ["单词书 CRUD", "单词增删改查", "公开 / 私有"])
    add_card(s3, Inches(8.8), Inches(1.5), Inches(3.8), Inches(2.2), "背单词", ["艾宾浩斯复习", "今日待复习", "学习进度"])
    add_card(s3, Inches(0.6), Inches(4.0), Inches(3.8), Inches(2.2), "单词自测", ["四选一出题", "自动计分", "更新复习进度"])
    add_card(s3, Inches(4.7), Inches(4.0), Inches(3.8), Inches(2.2), "错词统计", ["错词自动记录", "错误次数统计", "薄弱词分析"])
    add_card(s3, Inches(8.8), Inches(4.0), Inches(3.8), Inches(2.2), "考试记录", ["自测历史", "成绩查询", "答题明细"])
    add_footer(s3)

    # Slide 4 - Tech stack
    s4 = prs.slides.add_slide(blank)
    set_slide_bg(s4, WARM_WHITE)
    add_top_bar(s4, prs)
    add_title(s4, "03  技术架构")
    add_body(s4, [
        "前端：HTML + CSS + JavaScript（绿色暖白主题，单页应用）",
        "后端：Spring Boot 3.2 + Spring Data JPA",
        "数据库：H2 内存数据库（开发环境）",
        "通信：RESTful API + JSON 统一响应格式",
        "鉴权：Token（Authorization: Bearer xxx）",
        "部署：本地端口 3500，mvn spring-boot:run 启动",
    ], size=22)
    add_footer(s4)

    # Slide 5 - Architecture diagram text
    s5 = prs.slides.add_slide(blank)
    set_slide_bg(s5, WARM_WHITE)
    add_top_bar(s5, prs)
    add_title(s5, "04  系统分层设计")
    layers = [
        ("Controller 控制层", "接收 HTTP 请求，返回 JSON"),
        ("Service 业务层", "背单词逻辑、出题、错词记录"),
        ("Repository 数据层", "JPA 接口，操作数据库"),
        ("Entity 实体层", "User、Word、WordBook 等 7 张表"),
    ]
    y = Inches(1.55)
    for title, desc in layers:
        box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.5), y, Inches(8.3), Inches(0.95))
        box.fill.solid()
        box.fill.fore_color.rgb = GREEN_PALE if layers.index((title, desc)) % 2 == 0 else WHITE
        box.line.color.rgb = GREEN_LIGHT
        tb = s5.shapes.add_textbox(Inches(2.7), y + Inches(0.12), Inches(8), Inches(0.7))
        p = tb.text_frame.paragraphs[0]
        p.text = f"{title}  —  {desc}"
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT_DARK
        y += Inches(1.15)
    add_footer(s5)

    # Slide 6 - Database
    s6 = prs.slides.add_slide(blank)
    set_slide_bg(s6, WARM_WHITE)
    add_top_bar(s6, prs)
    add_title(s6, "05  数据库设计")
    add_body(s6, [
        "users — 用户账号信息",
        "word_books — 单词书（名称、描述、是否公开）",
        "words — 单词（拼写、音标、释义、例句）",
        "user_word_progress — 复习进度（阶段、下次复习时间）",
        "wrong_word_records — 错词记录（错误次数）",
        "exam_records / exam_answers — 考试记录与答题明细",
    ], size=21)
    add_footer(s6)

    # Slide 7 - Ebbinghaus highlight
    s7 = prs.slides.add_slide(blank)
    set_slide_bg(s7, WARM_WHITE)
    add_top_bar(s7, prs)
    add_title(s7, "06  核心亮点：艾宾浩斯间隔复习")
    highlight = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.5), Inches(11.9), Inches(2.0))
    highlight.fill.solid()
    highlight.fill.fore_color.rgb = GREEN_PALE
    highlight.line.color.rgb = GREEN_MID
    ht = s7.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(11.3), Inches(1.6))
    hp = ht.text_frame.paragraphs[0]
    hp.text = "5分钟 → 30分钟 → 12小时 → 1天 → 2天 → 4天 → 7天 → 15天 → 30天"
    hp.font.size = Pt(24)
    hp.font.bold = True
    hp.font.color.rgb = GREEN_DARK
    hp.alignment = PP_ALIGN.CENTER
    add_body(s7, [
        "答对：进入下一复习阶段，间隔自动延长",
        "答错：重置到第一阶段，并记入错词本",
        "实现文件：SpacedRepetitionUtil.java",
        "熟悉度根据阶段、正确/错误次数综合计算",
    ], top=Inches(3.8), size=21)
    add_footer(s7)

    # Slide 8 - API & integration
    s8 = prs.slides.add_slide(blank)
    set_slide_bg(s8, WARM_WHITE)
    add_top_bar(s8, prs)
    add_title(s8, "07  前后端联调与接口")
    add_body(s8, [
        "统一响应：{ code, message, data, timestamp }",
        "登录获取 Token → 后续请求携带 Authorization 头",
        "Apifox 管理接口文档，记录成功/失败/异常案例",
        "主要接口：/api/auth、/api/wordbooks、/api/memorize",
        "                  /api/quiz、/api/wrong-words、/api/exams",
        "跨域配置：允许所有来源，支持前后端分离开发",
    ], size=20)
    add_footer(s8)

    # Slide 9 - Demo flow
    s9 = prs.slides.add_slide(blank)
    set_slide_bg(s9, WARM_WHITE)
    add_top_bar(s9, prs)
    add_title(s9, "08  演示流程")
    steps = ["启动服务 mvn spring-boot:run", "登录 demo / password123", "首页查看学习统计", "单词书管理 / 添加单词", "背单词卡片复习", "四选一自测并查看分数", "错词统计 & 考试记录"]
    x = Inches(0.55)
    for i, step in enumerate(steps):
        col = i % 4
        row = i // 4
        left = Inches(0.55) + col * Inches(3.15)
        top = Inches(1.55) + row * Inches(2.5)
        circle = s9.shapes.add_shape(MSO_SHAPE.OVAL, left, top, Inches(0.55), Inches(0.55))
        circle.fill.solid()
        circle.fill.fore_color.rgb = GREEN_MID
        circle.line.fill.background()
        ct = circle.text_frame.paragraphs[0]
        ct.text = str(i + 1)
        ct.font.size = Pt(16)
        ct.font.bold = True
        ct.font.color.rgb = WHITE
        ct.alignment = PP_ALIGN.CENTER
        st = s9.shapes.add_textbox(left + Inches(0.65), top, Inches(2.4), Inches(1.8))
        sp = st.text_frame.paragraphs[0]
        sp.text = step
        sp.font.size = Pt(16)
        sp.font.color.rgb = TEXT_DARK
        sp.word_wrap = True
    add_footer(s9)

    # Slide 10 - Git & Summary
    s10 = prs.slides.add_slide(blank)
    set_slide_bg(s10, WARM_WHITE)
    add_top_bar(s10, prs)
    add_title(s10, "09  Git 管理与总结")
    add_body(s10, [
        "分支：main（正式版）/ develop（开发主线）/ feature-xxx",
        "提交规范：feat(auth): 实现用户登录功能",
        "项目地址：github.com/Nancy985/word-memory",
        "",
        "总结：完整前后端 + 科学记忆算法 + 错词追踪 + 自测评分",
        "展望：接入 MySQL、移动端适配、语音朗读、单词导入",
    ], size=21)
    thanks = s10.shapes.add_textbox(Inches(1), Inches(6.2), Inches(11.3), Inches(0.8))
    tp = thanks.text_frame.paragraphs[0]
    tp.text = "感谢聆听 · 欢迎提问"
    tp.font.size = Pt(28)
    tp.font.bold = True
    tp.font.color.rgb = GREEN_MID
    tp.alignment = PP_ALIGN.CENTER
    add_footer(s10, "")

    out = "/Users/mac/Projects/word-memory-backend/docs/词忆-答辩PPT.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
